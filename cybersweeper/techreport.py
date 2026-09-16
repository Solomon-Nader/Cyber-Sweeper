# Technical report export - a finished document (PDF or Word) for the user.
#
# Where report renders quick views (console text, Markdown,
# HTML), this module writes a *document*: cover page, executive summary,
# methodology, network context, host inventory, per-host detail, a findings
# register and a prioritised remediation plan, plus appendices. It is what a
# network owner would hand to a manager, or a student would attach to an
# assignment.
#
# Both output formats are built from the same ReportContent so they
# always say the same thing:
#
# * PDF - via reportlab (pure Python).
# * DOCX - via python-docx (pure Python), editable in Word/LibreOffice.
#
# Both libraries are regular dependencies of Cyber Sweeper; should either be
# missing, a clear RuntimeError tells the user what to install.

from __future__ import annotations

import html
import re
from collections import Counter, OrderedDict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from cybersweeper import __version__, config
from cybersweeper.models import SEVERITY_ORDER, Host, Port, ScanResult, Vulnerability
from cybersweeper.report import ReportFilter
from cybersweeper.utils import human_duration, is_precise_version, truncate_words

# UNKNOWN is a real outcome - an NVD entry that carries no CVSS score - and it
# gets its own row. Folding it into INFO made section 1 disagree with section 5
# of the same document.
SEVERITIES = ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO", "UNKNOWN"]
RANKED = ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO"]

# How far a finding can be trusted.
OBSERVED = "observed"            # a rule fired on a service state seen directly
VERSION_MATCH = "version match"  # CVE matched against a precise product version
UNCONFIRMED = "unconfirmed"      # CVE matched with no precise version to match on
CONFIRMED = (OBSERVED, VERSION_MATCH)

SEVERITY_RGB = {"CRITICAL": (0x7F, 0x1D, 0x1D), "HIGH": (0xC2, 0x41, 0x0C), "MEDIUM": (0xB4, 0x53, 0x09),
                "LOW": (0x25, 0x63, 0xEB), "INFO": (0x4B, 0x55, 0x63), "UNKNOWN": (0x6B, 0x72, 0x80), "": (0x9A, 0xA5, 0xB1)}
NAVY = (0x0B, 0x3D, 0x5C)
LIGHT = (0xF0, 0xF4, 0xF8)
GREY_TEXT = (0x62, 0x7D, 0x98)

SEVERITY_MEANING = {
    "CRITICAL": "Exploitable remotely with severe impact; fix immediately.",
    "HIGH": "Serious weakness likely to be exploited; fix within days.",
    "MEDIUM": "Meaningful weakness; fix in the next maintenance window.",
    "LOW": "Minor issue or hardening opportunity.",
    "INFO": "Informational; no direct risk but worth knowing.",
    "UNKNOWN": "No CVSS score published; judge it from the description.",
}
DEFAULT_MEANING = "Severity not classified."


# --------------------------------------------------------------------------- #
# Content model (format independent)
# --------------------------------------------------------------------------- #


@dataclass
class HostSection:
    host: Host
    port_rows: list[list[str]]  # Port, State, Service, Product/version, Risk
    findings: list[dict[str, Any]]


@dataclass
class ReportContent:
    title: str
    subtitle: str
    generated_at: datetime
    meta: OrderedDict[str, str]
    overall_risk: str
    severity_counts: OrderedDict[str, int]           # every finding
    severity_confirmed: OrderedDict[str, int]        # observed + version match
    severity_unconfirmed: OrderedDict[str, int]      # matched on a version range
    unconfirmed_risk: str                            # highest severity among those
    summary_paragraphs: list[str]
    top_findings: list[dict[str, Any]]
    methodology: list[str]
    inventory_rows: list[list[str]]  # IP, Hostname, MAC / vendor, Open ports, Highest risk
    hosts: list[HostSection]
    findings_rows: list[dict[str, Any]]  # sorted register
    remediation: list[dict[str, Any]]  # prioritised unique actions
    notes: list[str]
    options: dict[str, Any]
    glossary: list[tuple[str, str]] = field(default_factory=list)


def _risk_of(host: Host) -> str:
    return host.max_severity or ""


# The inventory's "Highest risk" column, rated on confirmed findings only so it
# agrees with the overall rating on the cover.
def _host_risk(ip: str, confirmed: list[dict[str, Any]]) -> str:
    sevs = [r["severity"] for r in confirmed if r["host"] == ip]
    if not sevs:
        return "none confirmed"
    return max(sevs, key=lambda sev: SEVERITY_ORDER.get(sev, 0))


def _overall(counts: dict[str, int]) -> str:
    for sev in RANKED:
        if counts.get(sev):
            return sev
    return "NONE"


# Rate the network on what the scan actually established.
#
# The overall rating used to be the single highest severity anywhere in the
# register, which let one speculative match set the headline: a home router was
# rated CRITICAL because a 2007 Samba CVE scored 10.0, matched against a version
# string of "3.X - 4.X" that the CVE almost certainly does not cover. The rating
# now comes from confirmed findings only - things the scan saw directly, or CVEs
# matched against a precise version - and the unconfirmed maximum is reported
# beside it rather than folded into it. Nothing is dropped; the two are simply
# no longer added together.
def _split_by_confidence(register: list[dict[str, Any]]) -> tuple[list, list]:
    confirmed = [r for r in register if r["confidence"] in CONFIRMED]
    return confirmed, [r for r in register if r["confidence"] == UNCONFIRMED]


def _counts_of(rows: list[dict[str, Any]]) -> OrderedDict:
    c = Counter(r["severity"] for r in rows)
    return OrderedDict((sev, c.get(sev, 0)) for sev in SEVERITIES)


# How much weight a finding carries.
#
# A rule finding states something the scan observed directly: port 23 answered,
# therefore Telnet is exposed. A CVE finding is a keyword match, and it is only
# as good as the version string it matched on. nmap frequently reports a range
# ("Samba smbd 3.X - 4.X") or an open-ended guess ("2.0.8 or later"); reducing
# that to a search term pulls in every CVE for the whole family, most of which
# will not apply to the build actually running. Saying so is the difference
# between a report and a list of things to worry about.
def _confidence(vuln: Vulnerability, port: Port) -> str:
    if vuln.source != "nvd":
        return OBSERVED
    return VERSION_MATCH if is_precise_version(port.version) else UNCONFIRMED


# Build the findings register: one entry per (host, finding), NOT per port.
#
# The same CVE on 139/tcp and 445/tcp of one host is one Samba daemon with one
# weakness. Listing it twice doubled the apparent size of the problem - a scan
# of a single home router reported 41 "findings" that were really 21 issues.
# Affected ports are kept in a list, so nothing is hidden by the grouping.
#
# Returns the register and the number of (port, finding) instances behind it.
def _build_register(hosts: list[Host]) -> tuple[list[dict[str, Any]], int]:
    grouped: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()
    instances = 0
    for h in hosts:
        for p in h.ports:
            for v in p.vulnerabilities:
                instances += 1
                port_label = f"{p.number}/{p.protocol}"
                entry = grouped.get((h.ip, v.cve_id))
                if entry is None:
                    grouped[(h.ip, v.cve_id)] = {
                        "severity": v.severity if v.severity in SEVERITIES else "UNKNOWN",
                        "id": v.cve_id, "host": h.ip, "hostname": h.hostname,
                        "ports": [port_label], "port": port_label,
                        "service": p.service_label if p.product else p.service,
                        "summary": v.summary, "recommendation": v.recommendation,
                        "cvss": v.cvss_score, "url": v.url, "source": v.source,
                        "confidence": _confidence(v, p),
                    }
                elif port_label not in entry["ports"]:
                    entry["ports"].append(port_label)
                    entry["port"] = ", ".join(entry["ports"])
    register = list(grouped.values())
    register.sort(key=lambda r: (-SEVERITY_ORDER.get(r["severity"], 0), -(r["cvss"] or 0), r["host"], r["ports"][0]))
    return register, instances


# How many addresses the scan actually swept.
#
# result.hosts holds only what answered, so deriving the figure from it produced
# "5 / 5" for a sweep of a whole /24 - which reads as though 5 addresses were
# examined and every one was live. Prefer the discovery note the engine wrote
# ("arp/nmap: 5/254 hosts up"), and fall back to expanding the target.
def _addresses_scanned(result: ScanResult) -> int:
    note = next((n for n in result.notes if n.startswith("discovery via")), "")
    m = re.search(r"/(\d+)\s+hosts up", note)
    if m:
        return int(m.group(1))
    try:
        from cybersweeper.utils import parse_targets

        return len(parse_targets(result.options.target, max_hosts=1 << 24))
    except Exception:  # pragma: no cover - unparseable target, e.g. a hostname list
        return len(result.hosts)


# Explain an empty port table honestly.
#
# "No ports matched the report filter" was printed even for a host where the
# scan simply found nothing open - which reads as if results were withheld.
def _no_ports_note(section: HostSection) -> str:
    if not section.host.ports:
        return "No open TCP ports were found on this host."
    return "No ports matched the report filter (the host has {} port{} recorded).".format(len(section.host.ports), "s" if len(section.host.ports) != 1 else "")


# Derive everything the renderers need from a ScanResult.
def build_content(result: ScanResult, flt: Optional[ReportFilter] = None,
                  title: str = "Network Security Assessment Report",
                  organisation: str = "") -> ReportContent:
    flt = flt or ReportFilter()
    hosts = flt.apply(result)
    now = datetime.now()

    open_ports = sum(len(h.open_ports) for h in hosts)

    # --- network context from notes -----------------------------------------
    network_ctx = next((n.split(":", 1)[1].strip() for n in result.notes if n.startswith("target auto-detected")), "")
    discovery_note = next((n for n in result.notes if n.startswith("discovery via")), "")

    register, instances = _build_register(hosts)

    # --- counts: one per register entry, so they agree with section 5 ---------
    confirmed_rows, unconfirmed_rows = _split_by_confidence(register)
    counts = _counts_of(register)
    severity_counts = OrderedDict((s, n) for s, n in counts.items() if n or s != "UNKNOWN")
    severity_confirmed = _counts_of(confirmed_rows)
    severity_unconfirmed = _counts_of(unconfirmed_rows)
    # The headline rating comes from confirmed findings only; the unconfirmed
    # maximum is reported next to it instead of being folded into it.
    overall = _overall(severity_confirmed)
    unconfirmed_risk = _overall(severity_unconfirmed)
    total_findings = sum(severity_counts.values())
    unconfirmed = len(unconfirmed_rows)

    meta: OrderedDict[str, str] = OrderedDict()
    if organisation:
        meta["Organisation"] = organisation
    meta["Target"] = result.options.target
    if network_ctx:
        meta["Network context"] = network_ctx
    meta["Scan started"] = result.started_at.strftime("%Y-%m-%d %H:%M:%S")
    meta["Duration"] = human_duration(result.duration)
    meta["Scan engine"] = result.engine or "n/a"
    meta["Ports examined"] = result.options.ports
    meta["Host discovery"] = discovery_note.replace("discovery via ", "") if discovery_note else (
        "not needed (single host)" if len(result.hosts) <= 1 else ("enabled" if result.options.discover else "disabled"))
    meta["CVE lookup"] = "NVD API 2.0 + built-in rules" if result.options.cve_lookup else "built-in rules only (offline)"
    # "5 / 5" invited the reading that only 5 addresses were examined. Report the
    # size of the range that was actually swept.
    scanned = _addresses_scanned(result)
    meta["Addresses scanned"] = str(scanned) if scanned else "n/a"
    meta["Hosts responding"] = str(len([h for h in result.hosts if h.is_up]))
    meta["Open ports"] = str(open_ports)
    meta["Findings"] = str(total_findings) + (f" ({instances} across all ports)" if instances != total_findings else "")
    meta["Overall risk rating"] = f"{overall} (from {len(confirmed_rows)} confirmed finding"
    meta["Overall risk rating"] += "s)" if len(confirmed_rows) != 1 else ")"
    if unconfirmed:
        meta["Unconfirmed CVE matches"] = (f"{unconfirmed} (highest {unconfirmed_risk}, not verified - "
                                           "excluded from the rating above)")
    meta["Tool"] = f"Cyber Sweeper {__version__}"
    if result.scan_id:
        meta["Scan id"] = f"#{result.scan_id}"

    # --- remediation plan: unique recommendations, highest severity first ------
    plan: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for r in register:
        key = r["recommendation"] or f"Review {r['id']}"
        entry = plan.setdefault(key, {"action": key, "severity": r["severity"], "targets": [], "ids": []})
        if SEVERITY_ORDER.get(r["severity"], 0) > SEVERITY_ORDER.get(entry["severity"], 0):
            entry["severity"] = r["severity"]
        # r["ports"] holds every port the grouped finding was seen on; splitting
        # the display string would silently drop all but the first.
        for port_label in r["ports"]:
            tgt = f"{r['host']}:{port_label.split('/')[0]}"
            if tgt not in entry["targets"]:
                entry["targets"].append(tgt)
        if r["id"] not in entry["ids"]:
            entry["ids"].append(r["id"])
    remediation = sorted(plan.values(), key=lambda e: -SEVERITY_ORDER.get(e["severity"], 0))
    for i, e in enumerate(remediation, start=1):
        e["priority"] = i

    # --- host sections ---------------------------------------------------------
    inventory, sections = [], []
    for h in hosts:
        inventory.append([h.ip, h.hostname or "-", " ".join(x for x in (h.mac, f"({h.vendor})" if h.vendor else "") if x) or "-",
                          str(len(h.open_ports)), _host_risk(h.ip, confirmed_rows)])
        port_rows = []
        for p in h.ports:
            risk = max((v.severity for v in p.vulnerabilities), key=lambda s: SEVERITY_ORDER.get(s, 0), default="")
            port_rows.append([f"{p.number}/{p.protocol}", p.state, p.service or "-", p.service_label if p.product else "-", risk or "-"])
        sections.append(HostSection(host=h, port_rows=port_rows,
                                    findings=[r for r in register if r["host"] == h.ip]))

    # --- executive summary text ----------------------------------------------------
    live = len(hosts)
    para1 = (f"Cyber Sweeper inspected {result.options.target}"
             + (f" ({network_ctx})" if network_ctx else "")
             + f" on {result.started_at:%d %B %Y at %H:%M}. {live} host{'s' if live != 1 else ''} "
             f"{'were' if live != 1 else 'was'} found live with {open_ports} open TCP port{'s' if open_ports != 1 else ''} "
             f"in total. The scan took {human_duration(result.duration)} using the {result.engine or 'socket'} engine.")
    if total_findings:
        n_conf = len(confirmed_rows)
        conf_parts = [f"{n} {s.lower()}" for s, n in severity_confirmed.items() if n]
        para2 = (f"{total_findings} distinct finding{'s' if total_findings != 1 else ''} "
                 f"{'were' if total_findings != 1 else 'was'} recorded"
                 + (f", affecting {instances} host-and-port combinations" if instances != total_findings else "")
                 + ". They fall into two groups, and the difference matters.\n\n"
                 + (f"{n_conf} {'are' if n_conf != 1 else 'is'} confirmed"
                    + (f" ({', '.join(conf_parts)})" if conf_parts else "")
                    + ": the scan either observed the condition directly - a service answering on a port it should "
                    "not - or matched a CVE against a precise product version. "
                    f"The overall risk rating for this network, {overall}, is based on these alone. "
                    if n_conf else
                    "No finding was confirmed: nothing was observed directly and no CVE matched a precise version. "
                    f"The overall risk rating is therefore {overall}. ")
                 + ("Critical and high confirmed findings should be addressed immediately; "
                    if severity_confirmed.get("CRITICAL") or severity_confirmed.get("HIGH") else "")
                 + "the remediation plan in section 6 lists the actions in priority order.")
        if unconfirmed:
            para2 += (f"\n\nThe remaining {unconfirmed} {'are' if unconfirmed != 1 else 'is'} unconfirmed"
                      f"{f', the most severe of which NVD scores as {unconfirmed_risk}' if unconfirmed_risk != 'NONE' else ''}. "
                      "These services did not report a precise version - nmap returned a range such as "
                      "'3.X - 4.X' or an open-ended '2.0.8 or later' - so each CVE was matched against the product "
                      "family instead of the build in use, and may well not apply to it. They are deliberately "
                      "excluded from the rating above, because one 10.0 score attached to a version that was never "
                      "identified would otherwise decide how this whole network is described. They are listed in "
                      "full in section 5 and should be checked against the vendor's advisory, or the version "
                      "confirmed on the device, before any of them is treated as real.")
    else:
        para2 = "No security findings were recorded for the open ports observed. Continue to monitor the network and re-scan after changes."
    para3 = ("Findings come from two sources: built-in configuration rules (identifiers CS-RULE-xxx), which flag "
             "services that are risky regardless of version, and the NIST National Vulnerability Database "
             "(identifiers CVE-yyyy-nnnn), matched on the software product and version identified on each port."
             if result.options.cve_lookup else
             "Findings come from Cyber Sweeper's built-in configuration rules (identifiers CS-RULE-xxx), which flag "
             "services that are risky regardless of version. CVE lookup against the National Vulnerability "
             "Database was not enabled for this scan; enabling it adds version-specific vulnerabilities.")

    methodology = [
        "Scope. The assessment covered the target range listed above. Only TCP services were examined. "
        "No exploitation, credential guessing or configuration changes were attempted; the tool's footprint "
        "on each target is limited to TCP connections, reading service greetings and, on web ports, a single HTTP HEAD request.",
        f"Host discovery. {('Live hosts were identified using ' + discovery_note.replace('discovery via ', '').split(':')[0] + ' probes.') if discovery_note else 'Discovery was not required for a single-host target.'} "
        "A host answering any probe - including a TCP reset from a closed port - is considered live.",
        f"Port scanning. {'nmap was used (arguments -sV -Pn -n) through the python-nmap library, giving port states and service versions from nmap probe database.' if result.engine == 'nmap' else 'Cyber Sweeper multithreaded socket engine performed TCP connect scans; open ports were fingerprinted from their banners.'} "
        f"Ports examined: {result.options.ports}. Port states are reported as open (service answered), closed (host refused) or filtered (no answer, usually a firewall).",
        "Service identification. Products and versions were taken from nmap -sV data where available, otherwise from "
        "banner fingerprints, falling back to the IANA well-known port table.",
        "Vulnerability assessment. Each open port was checked against 20 built-in rules for risky services "
        "(clear-text protocols, exposed remote administration, unauthenticated data stores). "
        + ("Identified products and versions were additionally searched in the NVD CVE database; CVSS base scores "
           "were mapped to severity bands (Critical 9.0-10, High 7.0-8.9, Medium 4.0-6.9, Low 0.1-3.9)."
           if result.options.cve_lookup else "NVD CVE lookup was not enabled."),
        "Counting. Each finding is counted once per host. A weakness in one service that listens on several ports "
        "appears as a single entry with its ports listed together, so the totals reflect the number of problems "
        "to fix rather than the number of open sockets they were seen on.",
        "Rating. The overall risk rating is derived from confirmed findings only. Unconfirmed CVE matches are reported in full, and their highest severity is stated beside the rating, but they do not set it: a single CVSS 10.0 attached to a product whose version was never identified would otherwise determine how an entire network is described.",
        "Confidence. Every finding carries a confidence value. 'observed' means the scan saw the condition "
        "directly and it needs no further proof. 'version match' means a CVE was matched against a precise "
        "product version. 'unconfirmed' means the service did not report a precise version - nmap gave a range "
        "such as '3.X - 4.X' or an open-ended '2.0.8 or later' - so the CVE was matched against the product "
        "family and may not apply to the build in use. Unconfirmed findings are evidence to check, not "
        "vulnerabilities that have been demonstrated.",
        "Limitations. UDP services were not scanned. Keyword-based CVE matching can include vulnerabilities that "
        "do not apply to the exact build in use, or miss some that do; every CVE listed should be verified against "
        "the vendor's advisory before action. CVE entries that NVD has withdrawn (marked 'Rejected') are excluded. "
        "Hosts protected by firewalls that drop probes may be under-reported. The tool reports exposure, and does "
        "not attempt to exploit anything, so no finding here is proof that a host can actually be compromised.",
    ]

    glossary = [
        ("Open / closed / filtered", "Port answered / host refused the connection / no answer (probably a firewall)."),
        ("CVE", "Common Vulnerabilities and Exposures - a public identifier for a known vulnerability (e.g. CVE-2023-38408)."),
        ("CVSS", "Common Vulnerability Scoring System - severity score from 0 to 10 published with most CVEs."),
        ("CS-RULE", "Cyber Sweeper built-in rule flagging a risky service configuration independent of version."),
        ("Banner", "The greeting text a service sends when a client connects; often reveals product and version."),
        ("SSID", "The name of a Wi-Fi network."),
        ("Confidence", "observed - the scan saw the condition directly (a rule finding). version match - the CVE was matched against a precise product version. unconfirmed - no precise version was available, so the CVE was matched against the product family and may not apply to the build in use."),
    ]

    subtitle = network_ctx or result.options.target
    return ReportContent(
        title=title, subtitle=subtitle, generated_at=now, meta=meta, overall_risk=overall,
        severity_counts=severity_counts, severity_confirmed=severity_confirmed,
        severity_unconfirmed=severity_unconfirmed, unconfirmed_risk=unconfirmed_risk,
        summary_paragraphs=[para1, para2, para3],
        top_findings=(confirmed_rows + unconfirmed_rows)[:6], methodology=methodology, inventory_rows=inventory, hosts=sections,
        findings_rows=register, remediation=remediation, notes=list(result.notes),
        options=result.options.to_dict(), glossary=glossary,
    )


# --------------------------------------------------------------------------- #
# PDF (reportlab)
# --------------------------------------------------------------------------- #


def _need(module: str, package: str):
    try:
        return __import__(module)
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise RuntimeError(
            f"{package} is required for this export. Install it with:  pip install {package}") from exc


def _esc(text: Any) -> str:
    return html.escape(str(text if text is not None else ""), quote=False)


# Write a PDF technical report and return its path.
def write_pdf(result: ScanResult, path: Path | str, flt: Optional[ReportFilter] = None,
              title: str = "Network Security Assessment Report", organisation: str = "") -> Path:
    _need("reportlab", "reportlab")
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    content = build_content(result, flt, title, organisation)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    def rgb(t):
        return colors.Color(t[0] / 255, t[1] / 255, t[2] / 255)

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=6)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, spaceAfter=0)
    cell = ParagraphStyle("cell", parent=body, fontSize=8, leading=10, spaceAfter=0)
    cell_b = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, textColor=rgb(NAVY), spaceBefore=10, spaceAfter=8)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, textColor=rgb(NAVY), spaceBefore=8, spaceAfter=4)
    cover_title = ParagraphStyle("ct", parent=body, fontName="Helvetica-Bold", fontSize=26, leading=32, textColor=rgb(NAVY), alignment=TA_CENTER)
    cover_sub = ParagraphStyle("cs", parent=body, fontSize=13, leading=17, textColor=rgb(GREY_TEXT), alignment=TA_CENTER)

    def table(data, col_widths, header=True, zebra=True, font_size=8):
        rows = [[Paragraph(_esc(c), cell_b if (header and i == 0) else cell) if not hasattr(c, "wrap") else c
                 for c in row] for i, row in enumerate(data)]
        t = Table(rows, colWidths=col_widths, repeatRows=1 if header else 0)
        style = [("GRID", (0, 0), (-1, -1), 0.4, rgb((0xC0, 0xD0, 0xE0))),
                 ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                 ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        if header:
            style.append(("BACKGROUND", (0, 0), (-1, 0), rgb(NAVY)))
        if zebra:
            for r in range(1 if header else 0, len(rows)):
                if r % 2 == 0:
                    style.append(("BACKGROUND", (0, r), (-1, r), rgb(LIGHT)))
        t.setStyle(TableStyle(style))
        return t

    def sev_chip(sev: str):
        p = ParagraphStyle("chip", parent=cell, fontName="Helvetica-Bold", textColor=colors.white, alignment=TA_CENTER)
        t = Table([[Paragraph(_esc(sev or "-"), p)]], colWidths=[1.9 * cm])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), rgb(SEVERITY_RGB.get(sev, SEVERITY_RGB[""]))),
                               ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        return t

    story: list = []
    # ---- cover -----------------------------------------------------------
    story += [Spacer(1, 4 * cm), Paragraph(_esc(content.title), cover_title), Spacer(1, 0.4 * cm),
              Paragraph(_esc(content.subtitle), cover_sub), Spacer(1, 1.2 * cm)]
    story.append(Table([[Paragraph(f"Overall risk rating: <b>{_esc(content.overall_risk)}</b>",
                                   ParagraphStyle("or", parent=body, fontSize=14, textColor=colors.white, alignment=TA_CENTER))]],
                       colWidths=[9 * cm], style=TableStyle([
                           ("BACKGROUND", (0, 0), (-1, -1), rgb(SEVERITY_RGB.get(content.overall_risk, SEVERITY_RGB["INFO"]))),
                           ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]), hAlign="CENTER"))
    story.append(Spacer(1, 1.2 * cm))
    meta_rows = [[k, v] for k, v in content.meta.items()]
    story.append(table(meta_rows, [5 * cm, 10 * cm], header=False))
    story += [Spacer(1, 1.5 * cm),
              Paragraph(f"Generated {content.generated_at:%Y-%m-%d %H:%M} by Cyber Sweeper {__version__}. "
                        f"{_esc(config.LEGAL_NOTICE)} This report is confidential to the network owner.", small),
              PageBreak()]

    # ---- 1 executive summary ---------------------------------------------
    story.append(Paragraph("1. Executive summary", h1))
    for ptxt in content.summary_paragraphs:
        story.append(Paragraph(_esc(ptxt), body))
    sev_table = [["Severity", "Confirmed", "Unconfirmed", "Meaning"]] + [
        [s, str(content.severity_confirmed.get(s, 0)), str(content.severity_unconfirmed.get(s, 0)),
         SEVERITY_MEANING.get(s, DEFAULT_MEANING)] for s in content.severity_counts]
    t = table(sev_table, [2.6 * cm, 2.1 * cm, 2.4 * cm, 7.9 * cm])
    for i, s in enumerate(content.severity_counts, start=1):
        t.setStyle(TableStyle([("TEXTCOLOR", (0, i), (0, i), rgb(SEVERITY_RGB[s])), ("FONTNAME", (0, i), (0, i), "Helvetica-Bold")]))
    story += [Spacer(1, 0.2 * cm), t, Spacer(1, 0.3 * cm)]
    if content.top_findings:
        story.append(Paragraph("Top findings", h2))
        rows = [["Severity", "Confidence", "Finding", "Host : port(s)", "Summary"]]
        for f in content.top_findings:
            rows.append([sev_chip(f["severity"]), f["confidence"], f["id"],
                         Paragraph(_esc(f"{f['host']}:{f['port']}"), cell),
                         truncate_words(f["summary"], 200)])
        story.append(table(rows, [2.1 * cm, 2.3 * cm, 2.9 * cm, 2.9 * cm, 4.8 * cm]))

    # ---- 2 methodology -----------------------------------------------------
    meth = [Paragraph("2. Scope and methodology", h1)]
    for ptxt in content.methodology:
        head, _, rest = ptxt.partition(". ")
        meth.append(Paragraph(f"<b>{_esc(head)}.</b> {_esc(rest)}", body))
    story.append(KeepTogether(meth[:2]))
    story.extend(meth[2:])

    # ---- 3 inventory -------------------------------------------------------
    inv = [Paragraph("3. Host inventory", h1)]
    if content.inventory_rows:
        inv.append(table([["IP address", "Hostname", "MAC / vendor", "Open ports", "Highest risk"]] + content.inventory_rows,
                         [3 * cm, 4 * cm, 4.5 * cm, 1.8 * cm, 2.2 * cm]))
    else:
        inv.append(Paragraph("No live hosts matched the report filter.", body))
    story.append(KeepTogether(inv))

    # ---- 4 per host ---------------------------------------------------------
    story.append(Paragraph("4. Detailed results per host", h1))
    for sec in content.hosts:
        h = sec.host
        label = h.ip + (f"  ({h.hostname})" if h.hostname else "")
        block = [Paragraph(_esc(label), h2)]
        extra = " · ".join(x for x in (h.mac, h.vendor, h.os_guess, f"discovered via {h.discovery_method}" if h.discovery_method else "") if x)
        if extra:
            block.append(Paragraph(_esc(extra), small))
        if sec.port_rows:
            block.append(table([["Port", "State", "Service", "Product / version", "Risk"]] + sec.port_rows,
                               [2.2 * cm, 1.8 * cm, 3 * cm, 6 * cm, 2 * cm]))
        else:
            block.append(Paragraph(_esc(_no_ports_note(sec)), body))
        story.append(KeepTogether(block))
        if sec.findings:
            rows = [["Severity", "Finding", "Port", "Summary and recommendation"]]
            for f in sec.findings:
                txt = f"<b>{_esc(f['summary'])}</b>"
                if f["recommendation"]:
                    txt += f"<br/><font color='#334e68'>Recommendation: {_esc(f['recommendation'])}</font>"
                tail = [] if f["cvss"] is None else [f"CVSS {f['cvss']}"]
                tail.append(f"confidence: {f['confidence']}")
                txt += f"<br/><font color='#627d98'>{_esc(' · '.join(tail))}</font>"
                rows.append([sev_chip(f["severity"]), f["id"], Paragraph(_esc(f["port"]), cell), Paragraph(txt, cell)])
            story += [Spacer(1, 0.15 * cm), table(rows, [2.4 * cm, 3.0 * cm, 2.4 * cm, 7.2 * cm]), Spacer(1, 0.3 * cm)]

    # ---- 5 findings register ------------------------------------------------
    reg = [Paragraph("5. Findings register", h1)]
    if content.findings_rows:
        rows = [["#", "Severity", "Finding", "Host : port(s)", "Service", "CVSS", "Confidence"]]
        for i, f in enumerate(content.findings_rows, start=1):
            rows.append([str(i), sev_chip(f["severity"]), f["id"],
                         Paragraph(_esc(f"{f['host']}:{f['port']}"), cell), f["service"] or "-",
                         "" if f["cvss"] is None else str(f["cvss"]), f["confidence"]])
        reg.append(table(rows, [0.8 * cm, 2.6 * cm, 3.0 * cm, 3.3 * cm, 2.9 * cm, 1.2 * cm, 2.2 * cm]))
    else:
        reg.append(Paragraph("No findings.", body))
    story.append(KeepTogether(reg) if len(content.findings_rows) <= 25 else reg[0])
    if len(content.findings_rows) > 25:
        story.append(reg[1])

    # ---- 6 remediation plan -------------------------------------------------
    rem = [Paragraph("6. Remediation plan", h1)]
    if content.remediation:
        rem.append(Paragraph("Actions are ordered by the highest severity they address. Completing the first items removes the most risk.", body))
        rows = [["Priority", "Severity", "Action", "Applies to"]]
        for e in content.remediation:
            rows.append([str(e["priority"]), sev_chip(e["severity"]), e["action"], ", ".join(e["targets"][:8]) + (" …" if len(e["targets"]) > 8 else "")])
        rem.append(table(rows, [1.5 * cm, 2.2 * cm, 7.3 * cm, 4 * cm]))
    else:
        rem.append(Paragraph("No remediation required based on this scan.", body))
    story.append(KeepTogether(rem) if len(content.remediation) <= 15 else rem[0])
    if len(content.remediation) > 15:
        story.extend(rem[1:])

    # ---- appendix ------------------------------------------------------------
    story.append(Paragraph("Appendix A - Scan configuration", h1))
    story.append(table([["Option", "Value"]] + [[k, str(v)] for k, v in content.options.items()], [5 * cm, 10 * cm]))
    if content.notes:
        story.append(Paragraph("Scan notes", h2))
        for n in content.notes:
            story.append(Paragraph("• " + _esc(n), body))
    story.append(Paragraph("Appendix B - Glossary", h1))
    story.append(table([["Term", "Meaning"]] + [[k, v] for k, v in content.glossary], [4 * cm, 11 * cm]))
    story.append(Paragraph("Appendix C - Disclaimer", h1))
    story.append(Paragraph(_esc(
        "This report was generated automatically by Cyber Sweeper from a point-in-time network scan. It reflects the "
        "state of the network at the time of the scan only. Vulnerability matches are based on service banners and "
        "public databases and must be verified before action. " + config.LEGAL_NOTICE), body))

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(rgb(GREY_TEXT))
        canvas.drawString(2 * cm, 1.2 * cm, f"{content.title} - {content.subtitle}"[:110])
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                            bottomMargin=2 * cm, title=content.title, author=f"Cyber Sweeper {__version__}",
                            subject=content.subtitle)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return path


# --------------------------------------------------------------------------- #
# DOCX (python-docx)
# --------------------------------------------------------------------------- #


# Write a Word (.docx) technical report and return its path.
def write_docx(result: ScanResult, path: Path | str, flt: Optional[ReportFilter] = None,
               title: str = "Network Security Assessment Report", organisation: str = "") -> Path:
    _need("docx", "python-docx")
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    content = build_content(result, flt, title, organisation)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for lvl, size in ((1, 16), (2, 12.5)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(*NAVY)

    def shade(cell_obj, rgb_tuple):
        tc_pr = cell_obj._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "{:02X}{:02X}{:02X}".format(*rgb_tuple))
        tc_pr.append(shd)

    def add_table(headers, rows, widths_cm, sev_col: Optional[int] = None, font_pt: float = 9):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, htxt in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(htxt)
            run.bold = True
            run.font.size = Pt(font_pt)
            run.font.color.rgb = RGBColor(255, 255, 255)
            shade(c, NAVY)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(val if val is not None else ""))
                run.font.size = Pt(font_pt)
                if sev_col is not None and i == sev_col and val in SEVERITY_RGB:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    shade(cells[i], SEVERITY_RGB[val])
                elif ri % 2 == 1:
                    shade(cells[i], LIGHT)
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
        doc.add_paragraph()
        return t

    # ---- cover -----------------------------------------------------------
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(content.title)
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor(*NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(content.subtitle)
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(*GREY_TEXT)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Overall risk rating: {content.overall_risk}")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(*SEVERITY_RGB.get(content.overall_risk, SEVERITY_RGB["INFO"]))
    doc.add_paragraph()
    add_table(["Item", "Value"], [[k, v] for k, v in content.meta.items()], [5, 12])
    p = doc.add_paragraph()
    r = p.add_run(f"Generated {content.generated_at:%Y-%m-%d %H:%M} by Cyber Sweeper {__version__}. {config.LEGAL_NOTICE} "
                  "This report is confidential to the network owner.")
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(*GREY_TEXT)
    doc.add_page_break()

    # ---- 1 executive summary ---------------------------------------------
    doc.add_heading("1. Executive summary", level=1)
    for ptxt in content.summary_paragraphs:
        doc.add_paragraph(ptxt)
    add_table(["Severity", "Confirmed", "Unconfirmed", "Meaning"],
              [[s, content.severity_confirmed.get(s, 0), content.severity_unconfirmed.get(s, 0),
                SEVERITY_MEANING.get(s, DEFAULT_MEANING)] for s in content.severity_counts],
              [2.6, 2.2, 2.6, 9.6], sev_col=0)
    if content.top_findings:
        doc.add_heading("Top findings", level=2)
        add_table(["Severity", "Confidence", "Finding", "Host : port(s)", "Summary"],
                  [[f["severity"], f["confidence"], f["id"], f"{f['host']}:{f['port']}",
                    truncate_words(f["summary"], 200)] for f in content.top_findings],
                  [2.2, 2.4, 3.0, 3.2, 6.2], sev_col=0, font_pt=8.5)

    # ---- 2 methodology -----------------------------------------------------
    doc.add_heading("2. Scope and methodology", level=1)
    for ptxt in content.methodology:
        head, _, rest = ptxt.partition(". ")
        p = doc.add_paragraph()
        p.add_run(head + ". ").bold = True
        p.add_run(rest)

    # ---- 3 inventory -------------------------------------------------------
    doc.add_heading("3. Host inventory", level=1)
    if content.inventory_rows:
        add_table(["IP address", "Hostname", "MAC / vendor", "Open ports", "Highest risk"], content.inventory_rows,
                  [3.2, 4.2, 5, 2, 2.6], sev_col=4)
    else:
        doc.add_paragraph("No live hosts matched the report filter.")

    # ---- 4 per host ---------------------------------------------------------
    doc.add_heading("4. Detailed results per host", level=1)
    for sec in content.hosts:
        h = sec.host
        doc.add_heading(h.ip + (f"  ({h.hostname})" if h.hostname else ""), level=2)
        extra = " · ".join(x for x in (h.mac, h.vendor, h.os_guess, f"discovered via {h.discovery_method}" if h.discovery_method else "") if x)
        if extra:
            p = doc.add_paragraph()
            r = p.add_run(extra)
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(*GREY_TEXT)
        if sec.port_rows:
            add_table(["Port", "State", "Service", "Product / version", "Risk"], sec.port_rows, [2.4, 2, 3.2, 7, 2.4], sev_col=4)
        else:
            doc.add_paragraph(_no_ports_note(sec))
        if sec.findings:
            rows = []
            for f in sec.findings:
                txt = f["summary"]
                if f["recommendation"]:
                    txt += f"\nRecommendation: {f['recommendation']}"
                if f["cvss"] is not None:
                    txt += f"\nCVSS {f['cvss']}"
                txt += f"\nConfidence: {f['confidence']}"
                rows.append([f["severity"], f["id"], f["port"], txt])
            add_table(["Severity", "Finding", "Port", "Summary and recommendation"], rows, [2.4, 3.2, 2.6, 8.8], sev_col=0, font_pt=8.5)

    # ---- 5 register ------------------------------------------------------------
    doc.add_heading("5. Findings register", level=1)
    if content.findings_rows:
        add_table(["#", "Severity", "Finding", "Host : port(s)", "Service", "CVSS", "Confidence"],
                  [[i, f["severity"], f["id"], f"{f['host']}:{f['port']}", f["service"] or "-",
                    "" if f["cvss"] is None else f["cvss"], f["confidence"]]
                   for i, f in enumerate(content.findings_rows, start=1)],
                  [0.9, 2.2, 3.2, 3.6, 3.4, 1.3, 2.2], sev_col=1, font_pt=8.5)
    else:
        doc.add_paragraph("No findings.")

    # ---- 6 remediation ----------------------------------------------------------
    doc.add_heading("6. Remediation plan", level=1)
    if content.remediation:
        doc.add_paragraph("Actions are ordered by the highest severity they address. Completing the first items removes the most risk.")
        add_table(["Priority", "Severity", "Action", "Applies to"],
                  [[e["priority"], e["severity"], e["action"], ", ".join(e["targets"][:8]) + (" …" if len(e["targets"]) > 8 else "")]
                   for e in content.remediation], [1.6, 2.4, 8.4, 4.6], sev_col=1, font_pt=8.5)
    else:
        doc.add_paragraph("No remediation required based on this scan.")

    # ---- appendices ---------------------------------------------------------------
    doc.add_heading("Appendix A - Scan configuration", level=1)
    add_table(["Option", "Value"], [[k, v] for k, v in content.options.items()], [5, 12])
    if content.notes:
        doc.add_heading("Scan notes", level=2)
        for n in content.notes:
            doc.add_paragraph(n, style="List Bullet")
    doc.add_heading("Appendix B - Glossary", level=1)
    add_table(["Term", "Meaning"], [[k, v] for k, v in content.glossary], [4.5, 12.5])
    doc.add_heading("Appendix C - Disclaimer", level=1)
    doc.add_paragraph(
        "This report was generated automatically by Cyber Sweeper from a point-in-time network scan. It reflects the "
        "state of the network at the time of the scan only. Vulnerability matches are based on service banners and "
        "public databases and must be verified before action. " + config.LEGAL_NOTICE)

    # footer with page numbers
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"{content.title} · {content.subtitle} · Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(*GREY_TEXT)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r_el = OxmlElement("w:r")
    t_el = OxmlElement("w:t")
    t_el.text = "1"
    r_el.append(t_el)
    fld.append(r_el)
    footer._p.append(fld)

    doc.core_properties.title = content.title
    doc.core_properties.author = f"Cyber Sweeper {__version__}"
    doc.core_properties.subject = content.subtitle
    doc.save(str(path))
    return path


WRITERS = {"pdf": write_pdf, "docx": write_docx}


def write(result: ScanResult, fmt: str, path: Path | str, flt: Optional[ReportFilter] = None, **kwargs) -> Path:
    try:
        writer = WRITERS[fmt.lower()]
    except KeyError as exc:
        raise ValueError(f"unknown technical report format '{fmt}' (pdf, docx)") from exc
    return writer(result, path, flt, **kwargs)
